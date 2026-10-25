#!/usr/bin/env python3
"""Convert the 'Source Data' sheet of the Excel balance workbook into a Fish Balance Manager .fbd file.

Usage:  python xlsx_to_fbd.py BALANCE_PivotTable.xlsx [output.fbd] [--date YYYY-MM-DD]
Needs:  pip install openpyxl

Reads (sheet 'Source Data'):
  A Supplier, B Product, C Kgs, D Price      one entry per row (rows with no supplier/product are skipped)
  F book marker                              Cash1 / Book1 / Book2 ... marks where a book STARTS; it runs until the next
                                             marker. Names starting with 'Cash' become Cash books, everything else Debtor.
  G 'Debtor' / 'Cash' with amounts to the right (I, J, K ...) on the same row - one amount per paper book.
A Kgs typed as a sum or product (=40+30+20, =8*30, =(8*30)+16) is kept as typed.
"""
import re, sys, datetime
try:
    import openpyxl
except ImportError:
    sys.exit("This needs openpyxl: pip install openpyxl")

def fmt4(x): return f"{x:.4f}"

def main(argv):
    date = None; args = []
    i = 1
    while i < len(argv):
        if argv[i] == '--date': date = argv[i + 1]; i += 2
        else: args.append(argv[i]); i += 1
    if not args: sys.exit(__doc__)
    src = args[0]
    out = args[1] if len(args) > 1 else re.sub(r'\.xlsx?$', '', src, flags=re.I) + '.fbd'
    date = date or datetime.date.today().isoformat()
    wf = openpyxl.load_workbook(src)['Source Data']
    wv = openpyxl.load_workbook(src, data_only=True)['Source Data']

    debtor, cash = [], []
    for r in range(1, wf.max_row + 1):
        label = str(wf.cell(r, 7).value or '').strip().lower()
        if label in ('debtor', 'cash'):
            for c in range(9, wf.max_column + 1):  # I onwards
                f = wf.cell(r, c).value
                if f is None or f == '': continue
                text = str(f).lstrip('=').replace(' ', '')
                if not re.fullmatch(r'[0-9.+\-]+', text): text = fmt4(float(wv.cell(r, c).value))
                (debtor if label == 'debtor' else cash).append(text)

    rows = []; book = ''; btype = ''; skipped = []
    for r in range(2, wf.max_row + 1):
        mark = wf.cell(r, 6).value
        if mark is not None and str(mark).strip():
            book = str(mark).strip(); btype = 'C' if book.lower().startswith('cash') else 'D'
        sup, prod = wf.cell(r, 1).value, wf.cell(r, 2).value
        if not sup and not prod: continue
        kv, pv = wv.cell(r, 3).value, wv.cell(r, 4).value
        if not sup or not prod or kv is None or pv is None:
            skipped.append(r); continue
        kf = wf.cell(r, 3).value
        expr = ''
        if isinstance(kf, str) and kf.startswith('='):
            t = kf[1:].replace(' ', '')
            # sums, products and brackets of plain numbers are kept as typed (the app's Kgs box takes + - * and ( ))
            if re.fullmatch(r'[0-9.+\-*()]+', t) and re.search(r'[+\-*]', t[1:]) and t.count('(') == t.count(')'): expr = t
        sup, prod = str(sup).strip().replace('|', '/'), str(prod).strip().replace('|', '/')
        fields = [sup, prod, fmt4(float(kv)), fmt4(float(pv)), date, '', '0']
        if expr or book: fields.append(expr)
        if book: fields += [book, btype]
        rows.append('|'.join(fields))

    lines = ['DEBTOR=' + ';'.join(debtor), 'CASH=' + ';'.join(cash), 'DRAFT_SUPPLIER=', 'DRAFT_SPECIES=', 'DRAFT_KGS=',
             'DRAFT_PRICE=', 'DRAFT_NOTES=', 'DRAFT_DATE=' + date]
    if book: lines.append(f'DRAFT_BOOK={book}|{btype}')
    lines += ['BEGIN'] + rows + ['END']
    with open(out, 'w', encoding='utf-8', newline='\r\n') as fh: fh.write('\n'.join(lines) + '\n')
    print(f"{len(rows)} entries, {len(debtor)} Debtor and {len(cash)} Cash amounts -> {out}")
    if skipped: print("Skipped rows with missing data:", skipped)

if __name__ == '__main__': main(sys.argv)
